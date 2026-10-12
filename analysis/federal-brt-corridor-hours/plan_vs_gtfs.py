"""BRT running hours: Plan spreadsheet vs Build GTFS, with and without the late-night trips.

usage: plan_vs_gtfs.py <build_feed_dir> <Federal_Blvd_Service_Plan.xlsx> <out_csv>
Needs openpyxl (pip install openpyxl).

- Running hours = first to last timepoint (spreadsheet) or first departure to last
  arrival (GTFS) per trip, both directions, no layover.
- "Late-night" = trips departing after the Build feed's last timed BRT departure
  (21:55). The Build feed carries these 16 weekday trips as placeholders with every
  time 00:00, so it has no running time for them.
- The spreadsheet stores times after midnight as dates (1900-01-01 00:24 = 24:24);
  they are read as next-day times, not dropped.
"""
import csv, sys, datetime

feed, xlsx, out = sys.argv[1:4]
CUT = 21 * 3600 + 55 * 60


def sec(v):
    if isinstance(v, datetime.datetime):  # after midnight: Excel day 1 = 1900-01-01
        return (v - datetime.datetime(1899, 12, 31)).total_seconds()
    if isinstance(v, datetime.time):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, datetime.timedelta):
        return v.total_seconds()
    return None


def h(s):
    a, b, c = s.strip().split(":")
    return int(a) * 3600 + int(b) * 60 + int(c)


import openpyxl
wb = openpyxl.load_workbook(xlsx, data_only=True)
plan = {}
for day, sheet in (("Weekday", "WK Times"), ("Saturday", "SA Times"), ("Sunday/Holiday", "SU Times")):
    rows = list(wb[sheet].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r[0] == "Block")
    trips = []
    for c0 in (2, 19):  # southbound, northbound timepoint blocks (11 timepoints each)
        for r in rows[hdr + 1:]:
            t = [sec(x) for x in r[c0:c0 + 11]]
            idx = [i for i, x in enumerate(t) if x is not None]
            if len(idx) >= 2:
                trips.append((t[idx[0]], t[idx[-1]], len(idx) == 11))
    plan[day] = trips

tr = {t["trip_id"] for t in csv.DictReader(open(feed + "/trips.txt", encoding="utf-8-sig")) if t["route_id"] == "FED 4C"}
span = {}
for r in csv.DictReader(open(feed + "/stop_times.txt", encoding="utf-8-sig")):
    if r["trip_id"] in tr:
        s = span.setdefault(r["trip_id"], [10 ** 9, 0])
        s[0] = min(s[0], h(r["departure_time"])); s[1] = max(s[1], h(r["arrival_time"]))
timed = [(a, b) for a, b in span.values() if not (a == 0 and b == 0)]
placeholders = len(span) - len(timed)


def summ(trips):
    return len(trips), round(sum((b - a) / 3600 for a, b, *_ in trips), 2)


rows = []
for day, trips in plan.items():
    late = [x for x in trips if x[0] > CUT]
    rows.append([day, "Plan spreadsheet", *summ(trips), *summ([x for x in trips if x[0] <= CUT]),
                 *summ(late), sum(1 for x in late if x[2]), "%02d:%02d" % divmod(max(a for a, *_ in trips) // 60, 60)])
pl = plan["Weekday"]
late_hrs = summ([x for x in pl if x[0] > CUT])[1]
g = summ(timed)
rows.insert(1, ["Weekday", "Build GTFS", g[0] + placeholders, "", g[0], g[1], placeholders, "",
                "", "%02d:%02d" % divmod(max(a for a, _ in timed) // 60, 60)])
rows.insert(2, ["Weekday", "Build GTFS + plan late-night times (estimate)", g[0] + placeholders,
                round(g[1] + late_hrs, 2), g[0], g[1], placeholders, late_hrs, "", ""])
# Weekend "GTFS modeled": the feed has no weekend BRT, so scale each weekend day's plan
# hours (late-night trips omitted, as the feed omits them) by the weekday feed/plan
# ratio on the same non-late trips.
ratio = g[1] / summ([x for x in pl if x[0] <= CUT])[1]
for day in ("Saturday", "Sunday/Holiday"):
    n, hrs = summ([x for x in plan[day] if x[0] <= CUT])
    i = next(k for k, r in enumerate(rows) if r[0] == day) + 1
    rows.insert(i, [day.split("/")[0] + " (GTFS modeled)", "Plan x weekday GTFS/plan ratio %.4f" % ratio,
                    "", "", n, round(hrs * ratio, 2), "", "", "", ""])

with open(out, "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["day_type", "source", "all_trips", "all_hrs", "excl_late_trips", "excl_late_hrs",
                "late_trips", "late_hrs", "late_full_corridor_trips", "last_departure"])
    w.writerows(rows)
for r in rows:
    print(r)
